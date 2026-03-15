"""
DOCX → HWPX 변환 모듈
python-docx로 Word 문서를 읽어 HWPX 구조 JSON으로 변환
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def parse_docx(file_path: str) -> dict:
    """DOCX 파일 → HWPX 구조 JSON 변환"""
    doc = Document(file_path)

    sections = []
    title = ""

    # 문단 처리
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            sections.append({"type": "paragraph", "content": ""})
            continue

        style = {}
        align = "LEFT"

        # 정렬
        if para.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            align = "CENTER"
        elif para.alignment == WD_ALIGN_PARAGRAPH.RIGHT:
            align = "RIGHT"

        # 문단 스타일에서 폰트 크기/볼드 추출
        if para.runs:
            run = para.runs[0]
            if run.bold:
                style["bold"] = True
            if run.font.size:
                style["font_size_pt"] = int(run.font.size.pt)
            if run.font.color and run.font.color.rgb:
                style["text_color"] = f"#{run.font.color.rgb}"

        # 제목 감지 (첫 번째 볼드+큰 텍스트)
        if not title and style.get("bold") and style.get("font_size_pt", 10) >= 14:
            title = text

        # 헤딩 스타일 감지
        if para.style and para.style.name.startswith("Heading"):
            style["bold"] = True
            level = para.style.name.replace("Heading ", "").replace("Heading", "")
            try:
                level_num = int(level) if level else 1
            except ValueError:
                level_num = 1
            size_map = {1: 18, 2: 16, 3: 14}
            style["font_size_pt"] = size_map.get(level_num, 12)
            if not title and level_num == 1:
                title = text

        sections.append({
            "type": "paragraph",
            "content": text,
            "align": align,
            "style": style,
        })

    # 테이블 처리
    for table in doc.tables:
        n_rows = len(table.rows)
        n_cols = len(table.columns)

        if n_rows == 0 or n_cols == 0:
            continue

        col_ratios = [round(100 / n_cols, 1)] * n_cols
        remainder = 100 - sum(col_ratios)
        col_ratios[-1] = round(col_ratios[-1] + remainder, 1)

        cells = []
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                cell_style = {"align": "LEFT"}

                # 첫 행 = 헤더
                if ri == 0:
                    cell_style["bold"] = True
                    cell_style["align"] = "CENTER"
                    cell_style["is_header"] = True

                # 셀 내 볼드 감지
                for para in cell.paragraphs:
                    if para.runs and para.runs[0].bold:
                        cell_style["bold"] = True

                cells.append({
                    "row": ri,
                    "col": ci,
                    "text": cell_text,
                    "style": cell_style,
                })

        sections.append({
            "type": "table",
            "table": {
                "rows": n_rows,
                "cols": n_cols,
                "col_widths_ratio": col_ratios,
                "cells": cells,
            }
        })

    if not title:
        title = "변환된 문서"

    return {
        "document": {"title": title, "page_width_hu": 42520},
        "sections": sections,
    }
